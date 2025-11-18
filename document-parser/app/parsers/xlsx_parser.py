import pandas as pd
from pathlib import Path
from typing import List, Dict, Any
from openpyxl import load_workbook
from app.models.schemas import DocumentMetadata, DocumentStructure, TableData, DocumentType

class XLSXParser:
    """Parse XLSX documents using pandas and openpyxl"""
    
    def parse(self, file_path: Path) -> DocumentStructure:
        """Parse XLSX file and extract structure"""
        
        # Extract metadata
        metadata = self._extract_metadata(file_path)
        
        # Extract sheets as tables
        tables = []
        paragraphs = []
        
        # Read all sheets
        excel_file = pd.ExcelFile(file_path)
        
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            # Convert DataFrame to TableData
            table_data = TableData(
                headers=df.columns.tolist(),
                rows=df.values.tolist()
            )
            tables.append(table_data)
            
            # Add sheet summary to paragraphs
            paragraphs.append(f"Sheet: {sheet_name} ({len(df)} rows, {len(df.columns)} columns)")
        
        # Create structure
        structure = DocumentStructure(
            title=metadata.filename,
            paragraphs=paragraphs,
            tables=tables,
            metadata=metadata
        )
        
        return structure
    
    def _extract_metadata(self, file_path: Path) -> DocumentMetadata:
        """Extract XLSX metadata"""
        wb = load_workbook(filename=str(file_path), read_only=True)
        props = wb.properties
        
        metadata = DocumentMetadata(
            filename=file_path.name,
            filesize=file_path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            document_type=DocumentType.XLSX,
            page_count=len(wb.sheetnames),
            author=props.creator,
            created_at=str(props.created) if props.created else None,
            modified_at=str(props.modified) if props.modified else None
        )
        
        wb.close()
        return metadata
    
    def extract_text(self, file_path: Path) -> str:
        """Extract plain text from XLSX"""
        excel_file = pd.ExcelFile(file_path)
        text_parts = []
        
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            text_parts.append(f"=== Sheet: {sheet_name} ===\n")
            text_parts.append(df.to_string())
            text_parts.append("\n\n")
        
        return "\n".join(text_parts)
