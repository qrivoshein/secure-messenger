from docx import Document
from pathlib import Path
from typing import List, Dict, Any
from app.models.schemas import DocumentMetadata, DocumentStructure, TableData, DocumentType

class DOCXParser:
    """Parse DOCX documents using python-docx"""
    
    def parse(self, file_path: Path) -> DocumentStructure:
        """Parse DOCX file and extract structure"""
        doc = Document(str(file_path))
        
        # Extract metadata
        metadata = self._extract_metadata(doc, file_path)
        
        # Extract content
        headings = []
        paragraphs = []
        tables = []
        
        for element in doc.element.body:
            # Extract paragraphs
            if element.tag.endswith('p'):
                para = doc.paragraphs[len(paragraphs)]
                text = para.text.strip()
                
                if text:
                    # Check if heading
                    if para.style.name.startswith('Heading'):
                        level = int(para.style.name.replace('Heading ', ''))
                        headings.append({
                            "level": level,
                            "text": text,
                            "index": len(paragraphs)
                        })
                    
                    paragraphs.append(text)
            
            # Extract tables
            elif element.tag.endswith('tbl'):
                table_index = len(tables)
                table = doc.tables[table_index]
                table_data = self._parse_table(table)
                tables.append(table_data)
        
        # Create structure
        structure = DocumentStructure(
            title=metadata.filename,
            headings=headings,
            paragraphs=paragraphs,
            tables=tables,
            metadata=metadata
        )
        
        return structure
    
    def _extract_metadata(self, doc: Document, file_path: Path) -> DocumentMetadata:
        """Extract DOCX metadata"""
        core_props = doc.core_properties
        
        return DocumentMetadata(
            filename=file_path.name,
            filesize=file_path.stat().st_size,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            document_type=DocumentType.DOCX,
            author=core_props.author,
            created_at=str(core_props.created) if core_props.created else None,
            modified_at=str(core_props.modified) if core_props.modified else None
        )
    
    def _parse_table(self, table) -> TableData:
        """Parse a DOCX table"""
        rows_data = []
        headers = []
        
        for i, row in enumerate(table.rows):
            row_data = [cell.text.strip() for cell in row.cells]
            
            if i == 0:
                headers = row_data
            else:
                rows_data.append(row_data)
        
        return TableData(
            headers=headers,
            rows=rows_data
        )
    
    def extract_text(self, file_path: Path) -> str:
        """Extract plain text from DOCX"""
        doc = Document(str(file_path))
        text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        return text
